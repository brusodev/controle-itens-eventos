#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Limpeza do estoque fantasma deixado pelo bug de reversao.

CONTEXTO
--------
O bug (reverter_baixa_estoque devolvia TODAS as SAIDAs historicas de uma O.S.
a cada edicao) encheu a tabela movimentacoes_estoque de ENTRADAs falsas. Como
a disponibilidade passou a ser derivada desse ledger (SAIDA - ENTRADA), o saldo
de dezenas de estoques ficou inflado -- estoque que nao existe.

CORRECAO DE DADOS
-----------------
A fonte de verdade e o CONTRATO consumido: as linhas reais das O.S. vigentes
(itens_ordem_servico). Este script reconstroi o ledger para que o consumo
liquido de cada estoque seja exatamente esse consumo real:

  - arquiva TODA a movimentacoes_estoque atual (tabela _bkp_<ts> + CSV);
  - substitui por UMA SAIDA limpa por linha de O.S. vigente;
  - realinha o cache quantidade_gasto = consumo real.

O.S. canceladas nao geram SAIDA (net 0). Itens realmente consumidos acima do
contratado (ver relatorio_estoque_excedido) CONTINUAM refletindo o consumo real
-- este script nao "conserta" contrato, so remove o fantasma.

SEGURANCA
---------
  - DRY-RUN por padrao: nao grava nada, so gera os arquivos antes/depois.
  - --aplicar: faz backup do arquivo .db ANTES de qualquer escrita.
  - Tudo roda numa unica transacao.

USO
---
    python3 scripts/diagnostico/limpar_estoque_fantasma.py            # dry-run
    python3 scripts/diagnostico/limpar_estoque_fantasma.py --aplicar  # grava
"""
import os
import sys
import csv
import shutil
import argparse
import sqlite3
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DB_PATH = os.path.join(BACKEND_DIR, 'instance', 'controle_itens.db')
OUT_DIR = os.path.join(BACKEND_DIR, 'instance', 'reconciliacao')


def conv(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(' ', '')
    if not t or t in ('__', 'None'):
        return 0.0
    neg = t.startswith('-')
    t = t.lstrip('-')
    if not t.replace(',', '').replace('.', ''):
        return 0.0
    try:
        n = float(t.replace('.', '').replace(',', '.'))
    except ValueError:
        return 0.0
    return -n if neg else n


def fmt_cache(v):
    """Grava no mesmo estilo de quantidade_inicial: numero puro."""
    v = max(0.0, v)
    return str(int(v)) if v == int(v) else '%.2f' % v


# -------------------------------------------------------------- leitura de estado
def estado_estoques(conn):
    """
    Retorna, por estoque: contratado, consumo real (O.S. vigentes) e o net atual
    do ledger. disp_antes usa o ledger (o que o sistema exibe HOJE); disp_depois
    usa o consumo real (o que passara a exibir apos a limpeza).
    """
    cur = conn.cursor()
    cur.execute("""
        SELECT ios.item_id, os.regiao_estoque, SUM(ios.quantidade_total)
        FROM itens_ordem_servico ios
        JOIN ordens_servico os ON os.id = ios.ordem_servico_id
        WHERE COALESCE(os.status, 'emitida') != 'cancelada'
        GROUP BY ios.item_id, os.regiao_estoque
    """)
    consumo = {(r[0], r[1]): float(r[2] or 0) for r in cur.fetchall()}

    cur.execute("""
        SELECT e.id, e.item_id, e.regiao_numero, e.quantidade_inicial, e.quantidade_gasto,
               i.descricao, i.item_codigo, i.unidade,
               COALESCE((SELECT SUM(CASE WHEN m.tipo='SAIDA' THEN m.quantidade
                                         ELSE -m.quantidade END)
                         FROM movimentacoes_estoque m
                         WHERE m.estoque_regional_id = e.id), 0) AS net
        FROM estoque_regional e LEFT JOIN itens i ON i.id = e.item_id
        ORDER BY i.descricao, e.regiao_numero
    """)
    linhas = []
    for eid, item, reg, ini, gasto, desc, cod, uni, net in cur.fetchall():
        contratado = conv(ini)
        real = consumo.get((item, reg), 0.0)
        net = float(net or 0)
        linhas.append({
            'estoque_id': eid, 'item_id': item, 'regiao': reg,
            'descricao': desc or '?', 'codigo': cod or '', 'unidade': uni or '',
            'contratado': contratado,
            'gasto_antes': net, 'disp_antes': contratado - net,
            'gasto_depois': real, 'disp_depois': contratado - real,
            'cache_antes': gasto,
        })
    return linhas


# -------------------------------------------------------------- geracao de arquivos
HDR_FILL = PatternFill('solid', fgColor='1F4E78')
HDR_FONT = Font(color='FFFFFF', bold=True)
CHG_FILL = PatternFill('solid', fgColor='FFF2CC')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _hdr(ws, n):
    for c in range(1, n + 1):
        cel = ws.cell(row=1, column=c)
        cel.fill = HDR_FILL
        cel.font = HDR_FONT
        cel.alignment = Alignment(horizontal='center', vertical='center')
        cel.border = BORDER


def _auto(ws):
    for col in ws.columns:
        w = max((len(str(c.value)) if c.value is not None else 0) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(w, 55)


def gerar_excel_antes_depois(linhas, caminho, modo):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Antes x Depois'
    ws.append(['Item', 'Codigo', 'Regiao', 'Contratado',
               'Gasto ANTES (ledger)', 'Disp. ANTES',
               'Gasto DEPOIS (real)', 'Disp. DEPOIS',
               'Variacao disponivel', 'Mudou?'])
    for l in linhas:
        delta = l['disp_depois'] - l['disp_antes']
        mudou = abs(delta) > 0.01
        ws.append([l['descricao'], l['codigo'], l['regiao'], l['contratado'],
                   l['gasto_antes'], l['disp_antes'],
                   l['gasto_depois'], l['disp_depois'],
                   delta, 'SIM' if mudou else ''])
    _hdr(ws, 10)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.border = BORDER
        for col in (4, 5, 6, 7, 8, 9):
            row[col - 1].number_format = '#,##0.00'
        if row[9].value == 'SIM':
            for c in row:
                c.fill = CHG_FILL
    _auto(ws)
    ws.freeze_panes = 'A2'

    # Resumo
    mudaram = [l for l in linhas if abs(l['disp_depois'] - l['disp_antes']) > 0.01]
    ws2 = wb.create_sheet('Resumo')
    resumo = [
        ['LIMPEZA DE ESTOQUE FANTASMA - ANTES x DEPOIS'],
        ['Gerado em', datetime.now().strftime('%d/%m/%Y %H:%M:%S')],
        ['Modo', modo],
        [''],
        ['Total de estoques', len(linhas)],
        ['Estoques com disponibilidade corrigida', len(mudaram)],
        ['Estoque fantasma removido (soma)',
         sum(l['disp_antes'] - l['disp_depois'] for l in mudaram)],
        [''],
        ['ANTES = disponibilidade derivada do ledger poluido (o que o sistema exibia)'],
        ['DEPOIS = contratado - consumo real das O.S. vigentes (a verdade)'],
    ]
    for r in resumo:
        ws2.append(r)
    ws2['A1'].font = Font(bold=True, size=12)
    ws2['B7'].number_format = '#,##0.00'
    _auto(ws2)
    wb.save(caminho)
    return len(mudaram)


def dump_movimentacoes_csv(conn, caminho):
    cur = conn.cursor()
    cur.execute("PRAGMA table_info(movimentacoes_estoque)")
    cols = [r[1] for r in cur.fetchall()]
    cur.execute("SELECT %s FROM movimentacoes_estoque ORDER BY id" % ','.join(cols))
    with open(caminho, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(cols)
        n = 0
        for row in cur:
            w.writerow(row)
            n += 1
    return n


# -------------------------------------------------------------- reconstrucao
def reconstruir_ledger(conn, ts):
    """Substitui a movimentacoes_estoque por SAIDAs limpas. Assume transacao aberta."""
    cur = conn.cursor()

    # 1) arquivar tabela atual dentro do banco
    tabela_bkp = 'movimentacoes_estoque_bkp_%s' % ts
    cur.execute('CREATE TABLE %s AS SELECT * FROM movimentacoes_estoque' % tabela_bkp)
    arquivadas = cur.execute('SELECT COUNT(*) FROM %s' % tabela_bkp).fetchone()[0]

    # 2) limpar
    cur.execute('DELETE FROM movimentacoes_estoque')

    # 3) uma SAIDA por linha de O.S. vigente
    agora = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    cur.execute("""
        SELECT ios.ordem_servico_id, ios.item_id, e.id, ios.quantidade_total, os.numero_os
        FROM itens_ordem_servico ios
        JOIN ordens_servico os ON os.id = ios.ordem_servico_id
        JOIN estoque_regional e
          ON e.item_id = ios.item_id AND e.regiao_numero = os.regiao_estoque
        WHERE COALESCE(os.status, 'emitida') != 'cancelada'
          AND COALESCE(ios.quantidade_total, 0) > 0
    """)
    linhas = cur.fetchall()
    criadas = 0
    for os_id, item_id, est_id, qtd, numero in linhas:
        cur.execute("""
            INSERT INTO movimentacoes_estoque
                (ordem_servico_id, item_id, estoque_regional_id, quantidade, tipo,
                 data_movimentacao, observacao)
            VALUES (?, ?, ?, ?, 'SAIDA', ?, ?)
        """, (os_id, item_id, est_id, float(qtd), agora,
              'Reconstrucao ledger %s - O.S. %s' % (ts, numero)))
        criadas += 1

    # 4) realinhar cache quantidade_gasto = consumo real por estoque
    cur.execute("""
        SELECT estoque_regional_id, SUM(quantidade)
        FROM movimentacoes_estoque WHERE tipo='SAIDA'
        GROUP BY estoque_regional_id
    """)
    reais = {r[0]: float(r[1] or 0) for r in cur.fetchall()}
    cur.execute('SELECT id FROM estoque_regional')
    for (eid,) in cur.fetchall():
        cur.execute('UPDATE estoque_regional SET quantidade_gasto=? WHERE id=?',
                    (fmt_cache(reais.get(eid, 0.0)), eid))

    return tabela_bkp, arquivadas, criadas


def main():
    p = argparse.ArgumentParser(description='Limpeza do estoque fantasma')
    p.add_argument('--aplicar', action='store_true', help='grava (padrao: dry-run)')
    p.add_argument('--db', default=DB_PATH)
    args = p.parse_args()

    if not os.path.exists(args.db):
        print('[ERRO] Banco nao encontrado:', args.db)
        return 1

    os.makedirs(OUT_DIR, exist_ok=True)
    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    modo = 'APLICAR (grava)' if args.aplicar else 'DRY-RUN (nao grava)'

    print('LIMPEZA DE ESTOQUE FANTASMA')
    print('Banco :', args.db)
    print('Modo  :', modo)
    print('Saida :', OUT_DIR)
    print('')

    # ---- ANTES (sempre) ----
    conn_ro = sqlite3.connect('file:%s?mode=ro' % args.db, uri=True)
    linhas = estado_estoques(conn_ro)
    csv_antes = os.path.join(OUT_DIR, 'movimentacoes_ANTES_%s.csv' % ts)
    n_antes = dump_movimentacoes_csv(conn_ro, csv_antes)
    conn_ro.close()
    print('Movimentacoes atuais exportadas :', n_antes, '->', os.path.basename(csv_antes))

    xlsx_ad = os.path.join(OUT_DIR, 'estoque_antes_depois_%s.xlsx' % ts)
    n_mudam = gerar_excel_antes_depois(linhas, xlsx_ad, modo)
    fantasma = sum(l['disp_antes'] - l['disp_depois']
                   for l in linhas if abs(l['disp_depois'] - l['disp_antes']) > 0.01)
    print('Antes x Depois (Excel)          :', os.path.basename(xlsx_ad))
    print('Estoques que serao corrigidos   :', n_mudam)
    print('Estoque fantasma a remover      : %.2f unidades (somadas)' % fantasma)

    if not args.aplicar:
        print('')
        print('DRY-RUN: nada foi gravado. Confira os arquivos acima.')
        print('Para aplicar: python3 %s --aplicar' % os.path.relpath(__file__, BACKEND_DIR))
        print('(sera feito backup do .db automaticamente antes de qualquer escrita)')
        return 0

    # ---- APLICAR ----
    bkp_db = os.path.join(OUT_DIR, 'controle_itens_pre_limpeza_%s.db' % ts)
    shutil.copy2(args.db, bkp_db)
    print('')
    print('Backup do banco criado          :', os.path.basename(bkp_db))

    conn = sqlite3.connect(args.db)
    try:
        conn.execute('BEGIN')
        tabela_bkp, arquivadas, criadas = reconstruir_ledger(conn, ts)
        conn.commit()
    except Exception:
        conn.rollback()
        conn.close()
        print('[ERRO] Falha na reconstrucao. Rollback feito, banco intacto.')
        raise
    print('Movimentacoes arquivadas em      :', tabela_bkp, '(%d linhas)' % arquivadas)
    print('SAIDAs limpas criadas            :', criadas)

    # ---- DEPOIS ----
    conn_ro = sqlite3.connect('file:%s?mode=ro' % args.db, uri=True)
    csv_depois = os.path.join(OUT_DIR, 'movimentacoes_DEPOIS_%s.csv' % ts)
    n_depois = dump_movimentacoes_csv(conn_ro, csv_depois)
    conn_ro.close()
    conn.close()
    print('Movimentacoes finais exportadas :', n_depois, '->', os.path.basename(csv_depois))
    print('')
    print('CONCLUIDO. Registros antes/depois em:', OUT_DIR)
    return 0


if __name__ == '__main__':
    sys.exit(main())
