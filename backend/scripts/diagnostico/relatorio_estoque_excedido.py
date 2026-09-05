#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera um relatorio Excel dos itens consumidos ACIMA do estoque contratado.

Fonte de verdade: as linhas reais das O.S. vigentes (itens_ordem_servico),
NAO o ledger de movimentacoes -- que ficou poluido pelo bug de reversao.

Consumo real por (item, regiao) = SUM(quantidade_total) das O.S. nao canceladas.
Excedente = consumo real - quantidade_inicial (contratado).

USO:
    python3 scripts/diagnostico/relatorio_estoque_excedido.py
    python3 scripts/diagnostico/relatorio_estoque_excedido.py --db <caminho> --out <arquivo.xlsx>

So leitura (read-only). Nao altera o banco.
"""
import os
import sys
import argparse
import sqlite3
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DB_PATH = os.path.join(BACKEND_DIR, 'instance', 'controle_itens.db')


def conv(v):
    """Le quantidade que pode estar em texto BR ('20.000', '2.700,00') ou numero."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(' ', '')
    if not t or t in ('__', 'None'):
        return 0.0
    neg = t.startswith('-')
    t = t.lstrip('-')
    if not t.replace(',', '').replace('.', ''):
        return 0.0
    try:
        n = float(t.replace('.', '').replace(',', '.'))
    except ValueError:
        return 0.0
    return -n if neg else n


HDR_FILL = PatternFill('solid', fgColor='1F4E78')
HDR_FONT = Font(color='FFFFFF', bold=True)
EXC_FILL = PatternFill('solid', fgColor='FCE4D6')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def estiliza_cabecalho(ws, ncols):
    for col in range(1, ncols + 1):
        c = ws.cell(row=1, column=col)
        c.fill = HDR_FILL
        c.font = HDR_FONT
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = BORDER


def auto_largura(ws):
    for col in ws.columns:
        largura = max((len(str(c.value)) if c.value is not None else 0) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(largura, 60)


def main():
    parser = argparse.ArgumentParser(description='Relatorio Excel de estoque excedido')
    parser.add_argument('--db', default=DB_PATH)
    parser.add_argument('--out', default=None)
    args = parser.parse_args()

    if not os.path.exists(args.db):
        print('[ERRO] Banco nao encontrado:', args.db)
        return 1

    ts = datetime.now().strftime('%Y%m%d_%H%M%S')
    out = args.out or os.path.join(BACKEND_DIR, 'instance', 'reconciliacao',
                                   'itens_estoque_excedido_%s.xlsx' % ts)
    os.makedirs(os.path.dirname(out), exist_ok=True)

    conn = sqlite3.connect('file:%s?mode=ro' % args.db, uri=True)
    cur = conn.cursor()

    # Consumo real por (item, regiao) das O.S. vigentes
    cur.execute("""
        SELECT ios.item_id, os.regiao_estoque, SUM(ios.quantidade_total)
        FROM itens_ordem_servico ios
        JOIN ordens_servico os ON os.id = ios.ordem_servico_id
        WHERE COALESCE(os.status, 'emitida') != 'cancelada'
        GROUP BY ios.item_id, os.regiao_estoque
    """)
    consumo = {(r[0], r[1]): float(r[2] or 0) for r in cur.fetchall()}

    # Contratado por estoque
    cur.execute("""
        SELECT e.item_id, e.regiao_numero, e.quantidade_inicial, i.descricao,
               i.item_codigo, i.unidade
        FROM estoque_regional e LEFT JOIN itens i ON i.id = e.item_id
    """)
    excedidos = []
    for item_id, regiao, inicial, descricao, codigo, unidade in cur.fetchall():
        usado = consumo.get((item_id, regiao), 0.0)
        contratado = conv(inicial)
        if usado > contratado + 0.01:
            excedidos.append({
                'item_id': item_id, 'regiao': regiao,
                'descricao': descricao or '?', 'codigo': codigo or '',
                'unidade': unidade or '', 'contratado': contratado,
                'consumido': usado, 'excedente': usado - contratado,
            })
    excedidos.sort(key=lambda x: -x['excedente'])

    wb = openpyxl.Workbook()

    # ---- Aba 1: Resumo ----
    ws = wb.active
    ws.title = 'Resumo'
    ws.append(['Item', 'Codigo', 'Regiao', 'Unidade',
               'Contratado', 'Consumido', 'Excedente', '% acima'])
    for e in excedidos:
        pct = (e['excedente'] / e['contratado'] * 100) if e['contratado'] else 0
        ws.append([e['descricao'], e['codigo'], e['regiao'], e['unidade'],
                   e['contratado'], e['consumido'], e['excedente'], round(pct, 2)])
    estiliza_cabecalho(ws, 8)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.border = BORDER
        for col in (5, 6, 7):
            row[col - 1].number_format = '#,##0.00'
        row[7].number_format = '0.00"%"'
        row[6].fill = EXC_FILL
    auto_largura(ws)
    ws.freeze_panes = 'A2'

    # ---- Aba 2: Detalhe por O.S. ----
    ws2 = wb.create_sheet('Detalhe por O.S.')
    ws2.append(['Item', 'Regiao', 'Numero O.S.', 'Modulo', 'Status',
                'Qtd na O.S.', 'Data emissao', 'Responsavel'])
    for e in excedidos:
        cur.execute("""
            SELECT os.numero_os, os.modulo, COALESCE(os.status, 'emitida'),
                   ios.quantidade_total, os.data_emissao_completa, os.responsavel
            FROM itens_ordem_servico ios
            JOIN ordens_servico os ON os.id = ios.ordem_servico_id
            WHERE ios.item_id = ? AND os.regiao_estoque = ?
              AND COALESCE(os.status, 'emitida') != 'cancelada'
            ORDER BY os.numero_os
        """, (e['item_id'], e['regiao']))
        for numero, modulo, status, qtd, data_em, resp in cur.fetchall():
            ws2.append([e['descricao'], e['regiao'], numero, modulo, status,
                        float(qtd or 0), data_em or '', resp or ''])
    estiliza_cabecalho(ws2, 8)
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
        for c in row:
            c.border = BORDER
        row[5].number_format = '#,##0.00'
    auto_largura(ws2)
    ws2.freeze_panes = 'A2'

    # ---- Aba 3: Metodologia ----
    ws3 = wb.create_sheet('Metodologia')
    linhas = [
        ['RELATORIO DE ITENS CONSUMIDOS ACIMA DO ESTOQUE'],
        ['Gerado em', datetime.now().strftime('%d/%m/%Y %H:%M:%S')],
        ['Banco', args.db],
        [''],
        ['Fonte de verdade', 'Linhas reais das O.S. vigentes (itens_ordem_servico)'],
        ['Consumo real', 'SUM(quantidade_total) das O.S. nao canceladas, por item e regiao'],
        ['Contratado', 'estoque_regional.quantidade_inicial'],
        ['Excedente', 'Consumo real - Contratado (apenas quando positivo)'],
        [''],
        ['OBS: o ledger de movimentacoes NAO foi usado aqui porque ficou'],
        ['poluido pelo bug de reversao. Este relatorio reflete o consumo'],
        ['contratual verdadeiro, independente daquela sujeira.'],
        [''],
        ['Total de itens excedidos', len(excedidos)],
    ]
    for l in linhas:
        ws3.append(l)
    ws3['A1'].font = Font(bold=True, size=12)
    auto_largura(ws3)

    wb.save(out)
    conn.close()

    print('Itens excedidos encontrados:', len(excedidos))
    for e in excedidos:
        print('  - %-32s reg %s: contratado %s, consumido %s, EXCEDENTE %s' % (
            e['descricao'][:32], e['regiao'],
            ('%.2f' % e['contratado']), ('%.2f' % e['consumido']),
            ('%.2f' % e['excedente'])))
    print('\nExcel gerado em:\n ', out)
    return 0


if __name__ == '__main__':
    sys.exit(main())
