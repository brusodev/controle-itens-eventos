"""Gera XLSX formatado dos itens do módulo Organização, Grupo 2 (Interior).

Correlaciona item -> preço unitário, agrupando por categoria.
Fonte: instance/controle_itens.db (estoque_regional região 2).
"""
import os
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB = os.path.join(BASE, 'instance', 'controle_itens.db')
OUT = os.path.join(os.path.dirname(BASE), 'docs',
                   'itens_organizacao_grupo2_interior_precos.xlsx')

CATEGORIA_LABEL = {
    'montagem_decoracao': 'Montagem e Decoração',
    'equipamento_informatica': 'Equipamento de Informática / Áudio-Visual',
    'recursos_humanos': 'Recursos Humanos',
    'material_grafico_expediente': 'Material Gráfico / Expediente',
}


def preco_para_float(valor):
    """Converte preço no formato pt-BR ('1.059,96') para float."""
    s = str(valor or '0').strip()
    if not s or s == '__':
        return 0.0
    try:
        return float(s.replace('.', '').replace(',', '.'))
    except ValueError:
        return 0.0


def main():
    con = sqlite3.connect(DB)
    rows = con.execute(
        """
        SELECT c.nome, i.item_codigo, i.descricao, i.unidade, e.preco
        FROM itens i
        JOIN categorias c ON c.id = i.categoria_id
        JOIN estoque_regional e ON e.item_id = i.id AND e.regiao_numero = 2
        WHERE c.modulo = 'organizacao'
        ORDER BY c.nome, CAST(i.item_codigo AS INTEGER)
        """
    ).fetchall()
    con.close()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Organização G2 Interior'

    # Paleta
    azul = '1F4E78'
    azul_claro = 'D9E1F2'
    cinza = 'F2F2F2'
    laranja = 'FCE4D6'

    borda = Border(*(Side(style='thin', color='BFBFBF'),) * 4)
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right = Alignment(horizontal='right', vertical='center')

    # --- Título ---
    ws.merge_cells('A1:E1')
    t = ws['A1']
    t.value = 'Itens Organização — Grupo 2 (Interior)'
    t.font = Font(bold=True, size=14, color='FFFFFF')
    t.fill = PatternFill('solid', fgColor=azul)
    t.alignment = center
    ws.row_dimensions[1].height = 26

    ws.merge_cells('A2:E2')
    s = ws['A2']
    s.value = 'Preço unitário por diária/unidade (R$) — fonte: estoque região 2'
    s.font = Font(italic=True, size=9, color='595959')
    s.alignment = center

    cabecalhos = ['Categoria', 'Cód.', 'Descrição', 'Unidade', 'Preço (R$)']
    header_row = 4
    for col, titulo in enumerate(cabecalhos, start=1):
        c = ws.cell(row=header_row, column=col, value=titulo)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=azul)
        c.alignment = center
        c.border = borda

    r = header_row + 1
    categoria_atual = None
    total_geral = 0.0
    sem_preco = 0

    for nome_cat, codigo, descricao, unidade, preco in rows:
        label = CATEGORIA_LABEL.get(nome_cat, nome_cat)
        # Linha separadora de categoria
        if label != categoria_atual:
            categoria_atual = label
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            cc = ws.cell(row=r, column=1, value=label)
            cc.font = Font(bold=True, color=azul)
            cc.fill = PatternFill('solid', fgColor=azul_claro)
            cc.alignment = left
            cc.border = borda
            for col in range(2, 6):
                ws.cell(row=r, column=col).border = borda
            r += 1

        preco_f = preco_para_float(preco)
        total_geral += preco_f
        if preco_f == 0:
            sem_preco += 1

        valores = ['', codigo, descricao, unidade, preco_f]
        for col, val in enumerate(valores, start=1):
            c = ws.cell(row=r, column=col, value=val)
            c.border = borda
            if col == 2:
                c.alignment = center
            elif col == 5:
                c.alignment = right
                c.number_format = 'R$ #,##0.00'
                if preco_f == 0:
                    c.fill = PatternFill('solid', fgColor=laranja)
                    c.font = Font(color='C55A11', italic=True)
            else:
                c.alignment = left
        if r % 2 == 0:
            for col in (2, 3, 4):
                cell = ws.cell(row=r, column=col)
                if cell.fill.fgColor.rgb in (None, '00000000'):
                    cell.fill = PatternFill('solid', fgColor=cinza)
        r += 1

    # --- Rodapé / totais ---
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    resumo = ws.cell(
        row=r, column=1,
        value=f'Total de itens: {len(rows)}  •  Sem preço cadastrado: {sem_preco}',
    )
    resumo.font = Font(bold=True)
    resumo.alignment = Alignment(horizontal='left', vertical='center')

    # Larguras
    larguras = {'A': 42, 'B': 8, 'C': 60, 'D': 18, 'E': 14}
    for col, w in larguras.items():
        ws.column_dimensions[col].width = w

    ws.freeze_panes = 'A5'
    ws.sheet_view.showGridLines = False

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f'Gerado: {OUT}')
    print(f'{len(rows)} itens • {sem_preco} sem preço')


if __name__ == '__main__':
    main()
